"""Company policy knowledge base -- retrieval over a folder of documents.

Drop .txt / .md / .pdf / .docx files into POLICY_DIR (default
backend/data/policy). The folder is re-read automatically whenever a file is
added, changed or removed (checked on every question; cheap). Documents are
split into ~900-character passages and searched with BM25, a keyword ranking
that needs no model download and works on any laptop. The chatbot's language
model turns a Hindi / Hinglish question into English keywords before
searching, and a small synonym table helps the no-model (basic) mode.
"""
from __future__ import annotations

import math
import re
import threading
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import structlog

logger = structlog.get_logger(__name__)

SUPPORTED = {".txt", ".md", ".pdf", ".docx"}
CHUNK_CHARS = 900
MAX_FILE_BYTES = 20 * 1024 * 1024

STOP = set(
    "a an the and or of to in on for is are was were be been it this that with as at by from what which who how "
    "when where do does can i my we our you your me please tell about kya hai hain ka ki ke ko me mein se aur "
    "batao bataiye hota hoti kaise kitna kitne kitni koi".split()
)
# Hindi / Hinglish words -> English words likely used in a policy document.
SYNONYMS = {
    "chutti": "leave holiday", "chhutti": "leave holiday", "chuti": "leave holiday", "avkash": "leave",
    "chhuttiyan": "leave holidays", "tankhwah": "salary pay", "tankha": "salary pay", "vetan": "salary pay",
    "pagar": "salary pay", "der": "late", "deri": "late", "hazri": "attendance", "haziri": "attendance",
    "upasthiti": "attendance", "samay": "time timing", "timing": "time hours", "shift": "shift",
    "overtime": "overtime ot", "ot": "overtime", "bonus": "bonus", "notice": "notice resignation",
    "istifa": "resignation", "resign": "resignation", "bimari": "sick medical", "beemar": "sick medical",
    "sick": "sick medical", "maternity": "maternity", "wfh": "work from home remote", "dress": "dress uniform",
}


def _tokens(text: str) -> list[str]:
    return [t for t in re.findall(r"\w+", text.lower()) if t not in STOP and len(t) > 1]


def _expand(tokens: list[str]) -> list[str]:
    out = list(tokens)
    for t in tokens:
        if t in SYNONYMS:
            out.extend(SYNONYMS[t].split())
    return out


@dataclass
class Chunk:
    source: str
    index: int
    text: str
    tokens: list[str] = field(default_factory=list)


def read_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in (".txt", ".md"):
        return path.read_text(encoding="utf-8", errors="replace")
    if ext == ".pdf":
        from pypdf import PdfReader

        return "\n\n".join((p.extract_text() or "") for p in PdfReader(str(path)).pages)
    if ext == ".docx":
        import docx

        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(parts)
    return ""


def split_chunks(text: str, size: int = CHUNK_CHARS) -> list[str]:
    """Paragraph-aware split; a heading line stays with the text below it."""
    paras = [p.strip() for p in re.split(r"\n\s*\n", text.replace("\r", "")) if p.strip()]
    chunks: list[str] = []
    cur = ""
    for p in paras:
        # a markdown heading (## Leave) starts a new passage, so each section
        # of a policy is retrieved on its own
        if cur and re.match(r"#{1,6}\s", p):
            chunks.append(cur)
            cur = ""
        while len(p) > size:  # one huge paragraph: hard split on a sentence / space
            cut = max(p.rfind(". ", 0, size), p.rfind(" ", 0, size))
            cut = cut + 1 if cut > size // 2 else size
            if cur:
                chunks.append(cur)
                cur = ""
            chunks.append(p[:cut].strip())
            p = p[cut:].strip()
        if cur and len(cur) + len(p) + 2 > size:
            chunks.append(cur)
            cur = p
        else:
            cur = f"{cur}\n\n{p}" if cur else p
    if cur:
        chunks.append(cur)
    return chunks


class PolicyIndex:
    """BM25 over all passages. Rebuilds itself when the folder changes."""

    k1, b = 1.5, 0.75

    def __init__(self, folder: Path) -> None:
        self.folder = folder
        self._sig: tuple[Any, ...] | None = None
        self._lock = threading.Lock()
        self.chunks: list[Chunk] = []
        self.files: list[str] = []
        self.errors: dict[str, str] = {}
        self._df: Counter[str] = Counter()
        self._avg = 1.0

    def _signature(self) -> tuple[Any, ...]:
        if not self.folder.is_dir():
            return ()
        items = []
        for p in sorted(self.folder.rglob("*")):
            if p.is_file() and p.suffix.lower() in SUPPORTED and not p.name.startswith(("~$", ".")):
                st = p.stat()
                items.append((str(p), st.st_mtime_ns, st.st_size))
        return tuple(items)

    def refresh(self, force: bool = False) -> None:
        sig = self._signature()
        if not force and sig == self._sig:
            return
        with self._lock:
            chunks: list[Chunk] = []
            files: list[str] = []
            errors: dict[str, str] = {}
            for path_s, _m, size in sig:
                path = Path(path_s)
                rel = str(path.relative_to(self.folder))
                if size > MAX_FILE_BYTES:
                    errors[rel] = "file larger than 20 MB, skipped"
                    continue
                try:
                    text = read_text(path)
                except Exception as exc:  # noqa: BLE001 - a bad file must not break the bot
                    errors[rel] = str(exc)[:200]
                    continue
                parts = split_chunks(text)
                if not parts:
                    errors[rel] = "no text found (scanned PDF? save it as text / docx)"
                    continue
                files.append(rel)
                for i, t in enumerate(parts):
                    chunks.append(Chunk(rel, i, t, _tokens(f"{Path(rel).stem} {t}")))
            df: Counter[str] = Counter()
            for c in chunks:
                df.update(set(c.tokens))
            self.chunks, self.files, self.errors, self._df = chunks, files, errors, df
            self._avg = (sum(len(c.tokens) for c in chunks) / len(chunks)) if chunks else 1.0
            self._sig = sig
            logger.info("policy_indexed", files=len(files), chunks=len(chunks), errors=len(errors))

    def search(self, query: str, k: int = 4) -> list[dict[str, Any]]:
        self.refresh()
        q = _expand(_tokens(query))
        if not q or not self.chunks:
            return []
        n = len(self.chunks)
        scored: list[tuple[float, Chunk]] = []
        for c in self.chunks:
            tf = Counter(c.tokens)
            dl = len(c.tokens) or 1
            s = 0.0
            for t in set(q):
                f = tf.get(t, 0)
                if not f:
                    continue
                idf = math.log(1 + (n - self._df[t] + 0.5) / (self._df[t] + 0.5))
                s += idf * f * (self.k1 + 1) / (f + self.k1 * (1 - self.b + self.b * dl / self._avg))
            if s > 0:
                scored.append((s, c))
        scored.sort(key=lambda x: -x[0])
        return [{"source": c.source, "passage": c.index + 1, "score": round(s, 2), "text": c.text} for s, c in scored[:k]]

    def status(self) -> dict[str, Any]:
        self.refresh()
        return {"folder": str(self.folder.resolve()), "exists": self.folder.is_dir(), "files": self.files,
                "chunks": len(self.chunks), "errors": self.errors}


_index: PolicyIndex | None = None


def get_index() -> PolicyIndex:
    global _index
    from app.config import get_settings

    folder = Path(get_settings().policy_dir)
    if _index is None or _index.folder != folder:
        _index = PolicyIndex(folder)
    return _index
